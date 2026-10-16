import os
from pathlib import Path
import calendar, csv, hashlib, json, math, sys
import openpyxl
ROOT = Path(os.environ.get('DJK_PROJECT_ROOT', '.'))
OUT = ROOT / 'outputs/baihe_monthly_indices_20260921'
OUT.mkdir(parents=True, exist_ok=True)
HIST = ROOT / '汉江流域逐月径流量数据（1950-2018年）-用户/汉江-白河站-逐月径流量表.xlsx'
RECENT = Path(os.path.join(os.environ.get('DJK_PROJECT_ROOT', '.'), 'external_inputs/desktop/丹江口/磨峪湾白河流量数据(1)/成果2024-2026'))
FILES = [HIST] + [RECENT / f for f in ['61801700-2024年白河站逐日平均流量表.xlsx', '61801700-2025年白河站逐日平均流量表.xlsx', '61801700-2026-白河逐日平均流量表.xlsx']]

def write_csv(name, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

def num(v):
    if v is None or str(v).strip() in ('', 'NA'):
        return None
    z = float(v)
    if not math.isfinite(z) or z < 0:
        raise ValueError(v)
    return z

def prepare():
    monthly = []
    daily = []
    audit = []
    for p in FILES:
        audit.append(dict(file=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
    w = openpyxl.load_workbook(HIST, read_only=True, data_only=True)
    for i, (y, m, q) in enumerate(list(w.worksheets[0].values)[1:], 2):
        monthly.append(dict(month=f'{int(y):04}-{int(m):02}', Q=num(q), source='historical_monthly', source_cell=f'C{i}', days_valid='', days_expected='', reported_Q=''))
    w.close()
    for p in FILES[1:]:
        w = openpyxl.load_workbook(p, read_only=True, data_only=True)
        a = list(w['Sheet1'].values)
        y = int(a[1][2])
        assert str(a[1][5]).strip() == '61801700'
        for m in range(1, 13):
            nd = calendar.monthrange(y, m)[1]
            vv = []
            for day in range(1, nd + 1):
                q = num(a[day + 3][m + 1])
                if q is not None:
                    vv.append(q)
                    daily.append(dict(date=f'{y:04}-{m:02}-{day:02}', Q=q, source=p.name, source_cell=f'{openpyxl.utils.get_column_letter(m + 2)}{day + 4}'))
            reported = num(a[35][m + 1])
            if vv:
                assert len(vv) == nd, 'Incomplete month: no mean imputation permitted'
                q = sum(vv) / nd
                assert reported is not None and abs(q - reported) <= max(0.51, 0.0011 * reported)
                monthly.append(dict(month=f'{y:04}-{m:02}', Q=q, source=f'daily_aggregated_{y}', source_cell=f'{openpyxl.utils.get_column_letter(m + 2)}5:{openpyxl.utils.get_column_letter(m + 2)}{nd + 4}', days_valid=len(vv), days_expected=nd, reported_Q=reported))
        if y == 2026:
            b = list(w['Sheet2'].values)[1:]
            second = {str(r[1])[:10]: num(r[2]) for r in b if num(r[2]) is not None}
            first = {d['date']: d['Q'] for d in daily if d['date'].startswith('2026')}
            assert first == second, '2026 sheets disagree'
        w.close()
    assert len({r['month'] for r in monthly}) == len(monthly)
    assert sum((r['Q'] is not None for r in monthly)) == 510
    assert len(daily) == 546
    write_csv('monthly_inputs.csv', monthly)
    write_csv('recent_daily_inputs.csv', daily)
    (OUT / 'source_hashes.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Verified: 492 historical months + 18 recent months; 546 daily values; 2026 sheets agree.')
if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    prepare()
